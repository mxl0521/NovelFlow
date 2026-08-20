from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = next(root.joinpath("docs").glob("*.md"))
target = root / "docs" / "NovelFlow-Technical-Architecture.docx"
markdown_target = root / "docs" / "NovelFlow-Technical-Architecture.md"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
for style_name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)

lines = source.read_text(encoding="utf-8").splitlines()
markdown_target.write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
in_code = False
code_lines = []
table_rows = []
fence = chr(96) * 3

def flush_table():
    global table_rows
    if not table_rows:
        return
    rows = table_rows
    table_rows = []
    if len(rows) < 2:
        for row in rows:
            doc.add_paragraph(row)
        return
    cells = [[cell.strip() for cell in row.strip("|").split("|")] for row in rows if "---" not in row]
    if not cells:
        return
    table = doc.add_table(rows=1, cols=len(cells[0]))
    table.style = "Table Grid"
    for idx, value in enumerate(cells[0]):
        table.rows[0].cells[idx].text = value
    for row in cells[1:]:
        cells_out = table.add_row().cells
        for idx, value in enumerate(row):
            if idx < len(cells_out):
                cells_out[idx].text = value

for line in lines:
    if line.startswith(fence):
        flush_table()
        if in_code:
            p = doc.add_paragraph()
            p.style = "No Spacing"
            for code_line in code_lines:
                run = p.add_run(code_line + "\n")
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                run.font.size = Pt(8.5)
            code_lines = []
            in_code = False
        else:
            in_code = True
        continue
    if in_code:
        code_lines.append(line)
        continue
    if line.startswith("|"):
        table_rows.append(line)
        continue
    flush_table()
    if not line.strip():
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line[2:].strip())
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
    elif re.match(r"^\d+\. ", line):
        doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
    elif line.startswith("- "):
        doc.add_paragraph(line[2:].strip(), style="List Bullet")
    elif line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(line[2:].strip())
        run.italic = True
    elif line == "---":
        doc.add_page_break()
    else:
        doc.add_paragraph(line)

flush_table()
doc.save(target)
print(target)
